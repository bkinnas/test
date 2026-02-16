"""Create a sample salary adjustment letter .docx template for testing."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("SALARY ADJUSTMENT NOTIFICATION")
run.bold = True
run.font.size = Pt(16)

doc.add_paragraph()

# Date
doc.add_paragraph("Date: {{date}}")
doc.add_paragraph()

# Recipient
doc.add_paragraph("To: {{employee_name}}")
doc.add_paragraph("Employee ID: {{employee_id}}")
doc.add_paragraph("Department: {{department}}")
doc.add_paragraph("Position: {{position}}")
doc.add_paragraph()

# Body
doc.add_paragraph(
    "Dear {{first_name}},"
)
doc.add_paragraph()
doc.add_paragraph(
    "We are pleased to inform you that your compensation has been adjusted "
    "effective {{effective_date}}. The details of your salary adjustment are as follows:"
)
doc.add_paragraph()

# Salary table
table = doc.add_table(rows=4, cols=2)
table.style = "Table Grid"
cells = table.rows[0].cells
cells[0].text = "Previous Salary"
cells[1].text = "${{previous_salary}}"
cells = table.rows[1].cells
cells[0].text = "New Salary"
cells[1].text = "${{new_salary}}"
cells = table.rows[2].cells
cells[0].text = "Raise Amount"
cells[1].text = "${{raise_amount}}"
cells = table.rows[3].cells
cells[0].text = "Raise Percentage"
cells[1].text = "{{raise_percentage}}"

doc.add_paragraph()
doc.add_paragraph("Reason: {{reason}}")
doc.add_paragraph()
doc.add_paragraph(
    "This adjustment reflects our recognition of your contributions and performance. "
    "If you have any questions regarding this change, please contact your manager, "
    "{{manager_name}}, or the HR department."
)
doc.add_paragraph()
doc.add_paragraph("Sincerely,")
doc.add_paragraph("{{company_name}}")
doc.add_paragraph("Human Resources Department")

doc.save("sample_salary_adjustment_template.docx")
print("Sample template created: sample_salary_adjustment_template.docx")
